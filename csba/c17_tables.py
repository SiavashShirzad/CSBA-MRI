import sys, os; sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from stats import *
from figstyle import relabel
from docx import Document
from docx.shared import Pt, Mm
from docx.enum.section import WD_ORIENT

OUT = ROOT/'manuscript'/'tables'; OUT.mkdir(parents=True, exist_ok=True)


def P_(p):
    if pd.isna(p): return ''
    return '<0.001' if p < .001 else (f'{p:.1g}' if p < .01 else f'{p:.2g}' if p < .1 else f'{p:.2f}')


def ci(b, lo, hi, d=1): return f'{b:.{d}f} ({lo:.{d}f} to {hi:.{d}f})'.replace('-', '−')


G = load_gaps(); P = G[G['role'] != 'retest'].copy()

# Table 1 (compact; full breakdown by source in Supplementary Table 1)
def grp(r):
    if r['group'] == 'HC':
        return {'dev': '1 Development', 'test': '2 Internal test'}.get(r['role'], '3 External test (JHU, MCW)')
    return {'JHU_abnormal': '4 JHU, abnormal MRI', 'DCM': '5 Degenerative cervical myelopathy (JHU, MCW)', 'Duke_abnormal': '6 CSpineSeg, abnormal MRI',
            'SG_compression': '7 Spine Generic, asymptomatic cord compression'}.get(r['group'], '8 Other patients*')
P['tg'] = P.apply(grp, axis=1)
def t1row(d): return pd.Series({'Participants': len(d), 'Age, y, mean (SD)': f"{d.age.mean():.0f} ({d.age.std():.0f})",
                                'Women, %': f"{100*(d.sex == 'F').sum()/max(d.sex.notna().sum(), 1):.0f}", '3D MRI, %': f"{100*(d.protocol == '3D').mean():.0f}"})
T1 = P.groupby('tg').apply(t1row)
T1.index = [i[2:] for i in T1.index]
T1 = pd.concat([pd.DataFrame([['', '', '', '']], index=['Healthy reference'], columns=T1.columns), T1.iloc[:3],
                pd.DataFrame([['', '', '', '']], index=['Patients'], columns=T1.columns), T1.iloc[3:]])
T1.loc['Total'] = t1row(P)
T1.index = [i if i in ('Healthy reference', 'Patients', 'Total') else '   ' + i for i in T1.index]; T1.index.name = 'Group'
OTHER = P[P['tg'].str.startswith('8')]['group'].value_counts().to_dict()

# Table 2
A = pd.read_csv(TAB/f'accuracy_{PRIMARY}.csv')
order = ['Cross-validated development', 'Internal test', 'External Center 1 controls', 'External Center 2 controls']
kinds = ['raw', 'calibrated', 'calibrated + local offset (leave-one-out)']
T2 = []
for s in order:
    for k in kinds:
        r = A[(A['set'] == s) & (A['kind'] == k)]
        if r.empty: continue
        r = r.iloc[0]
        T2.append({'Set': s.replace(' controls', ' (healthy)'), 'Estimate': {'raw': 'Network output', 'calibrated': 'Calibrated',
                   'calibrated + local offset (leave-one-out)': 'Calibrated + local offset'}[k], 'Participants': int(r['n']),
                   'MAE, y (95% CI)': ci(r['mae'], r['mae_lo'], r['mae_hi'], 2), 'r (95% CI)': ci(r['r'], r['r_lo'], r['r_hi'], 3) if pd.notna(r['r_lo']) else f"{r['r']:.3f}",
                   'Mean error, y': f"{r['bias']:+.2f}".replace('-', '−'), 'Slope of error on age': f"{r['slope']:+.3f}".replace('-', '−')})
T2 = pd.DataFrame(T2)

# Table 3
C = pd.read_csv(TAB/f'contrasts_{PRIMARY}.csv'); BC = pd.read_csv(TAB/'contrast_calibration_bootstrap.csv').set_index('contrast')
NAME = {'Center 1 abnormal vs controls': 'Center 1: abnormal MRI vs controls', 'DCM vs local controls (Centers 1-2)': 'DCM vs local controls (Centers 1 and 2)',
        'Duke abnormal vs Duke controls': 'CSpineSeg: abnormal vs healthy', 'Spine Generic compression vs controls': 'Spine Generic: asymptomatic compression vs healthy'}
T3 = []
for cn in NAME:
    d = C[C['contrast'] == cn].set_index('marker'); w = d.loc['Whole spine']
    row = {'Contrast': NAME[cn], 'Controls / cases': f"{int(w['n_ctrl'])} / {int(w['n_case'])}",
           'Age, controls': w['age_ctrl'], 'Age, cases': w['age_case'], 'Adjustment': w['formula'].split('~')[1].replace('grp + ', '').replace('C(center)', 'center').replace('C(site)', 'site').strip(),
           'Whole spine, y (95% CI)': ci(w['beta'], w['lo'], w['hi']), 'Holm P': P_(w['p_holm_primary']),
           'Calibration bootstrap 95% CI': f"{BC.loc[cn, 'lo']:.1f} to {BC.loc[cn, 'hi']:.1f}"}
    for m in ['Cord', 'CSF', 'Disc', 'Vertebral bone']:
        r = d.loc[m]; row[f'{m}, y'] = ci(r['beta'], r['lo'], r['hi']) + f" [{P_(r['p_holm_compartments'])}]"
    T3.append(row)
T3 = pd.DataFrame(T3)

# Supplementary: baseline clinical correlates in DCM
CL = pd.read_csv(TAB/'clinical_endpoints.csv').set_index('endpoint')
T4 = []
for k, lab, unit in [('Baseline mJOA', 'Baseline mJOA', 'points'),
                     ('Severity (mild/moderate/severe, ordinal)', 'Severity category (mild / moderate / severe)', 'OR, more severe category'),
                     ('Frailty (mFI-5)', 'Modified frailty index (mFI-5)', 'points'), ('Cord T2 hyperintensity', 'Cord T2 hyperintensity', 'OR'),
                     ('MSCC (AP diameter)', 'Maximum spinal cord compression (AP diameter)', '%'), ('Symptom duration (months)', 'Symptom duration', 'months')]:
    r = CL.loc[k]; T4.append({'Measure': lab, 'Participants': int(r['n']), 'Unit': unit, 'Estimate per 1 SD of age gap (95% CI)': ci(r['beta'], r['lo'], r['hi'], 2), 'P': P_(r['p'])})
T4 = pd.DataFrame(T4)

# supplementary
def g(f): return pd.read_csv(TAB/f)
def est(d, k=1): return [ci(a, b, c, k) for a, b, c in zip(d.beta, d.lo, d.hi)]
DS = {'Duke': 'CSpineSeg', 'SpineGeneric': 'Spine Generic', 'OpenNeuro': 'OpenNeuro', 'JHU_clinical': 'Center 1 clinical', 'JHU_DCM': 'Center 1 DCM study', 'MCW': 'Center 2 DCM study'}
CN = {'Center 1 abnormal vs controls': 'Center 1 abnormal MRI', 'DCM vs local controls (Centers 1-2)': 'DCM (Centers 1 and 2)',
      'Duke abnormal vs Duke controls': 'CSpineSeg abnormal', 'Spine Generic compression vs controls': 'Asymptomatic compression'}
SUP = {}
F = g('flow_exclusions.csv'); SUP['S1'] = ('Image and participant exclusions by data source',
    F.assign(dataset=F.dataset.map(lambda x: DS.get(x, x))).rename(columns={'dataset': 'Data source', 'exclusion': 'Status / reason', 'images': 'Images', 'participants': 'Participants'}), '')
R = g('regional_performance.csv'); R = R[R.region.isin(['spine', 'cord', 'csf', 'disc', 'bone', 'C2', 'C3', 'C4', 'C5', 'C6', 'C7', 'T1'])]
SUP['S2'] = ('Accuracy of tissue- and level-specific age estimates, internal test set',
    pd.DataFrame({'Region': R.region.map({'spine': 'Whole spine', 'cord': 'Cord', 'csf': 'CSF', 'disc': 'Intervertebral discs', 'bone': 'Vertebral bone'}).fillna(R.region + ' level'),
                  'r (95% CI)': [ci(a, b, c, 3) for a, b, c in zip(R.r, R.r_lo, R.r_hi)], 'MAE, y': R.mae.round(2)}), 'Network output before calibration.')
SG = g('accuracy_subgroups.csv')
SGn = SG.set.str.replace('dataset=', '').map(lambda x: DS.get(x, x)).str.replace('protocol=', 'Acquisition ').str.replace('sex=F', 'Women').str.replace('sex=M', 'Men') \
    .str.replace(r'age_band=\((\d+), (\d+)\]', lambda m: f'Age {int(m.group(1)) + 1}–{m.group(2)} y', regex=True)
SUP['S3'] = ('Accuracy by subgroup, healthy development (out-of-fold) and internal test participants',
    pd.DataFrame({'Subgroup': SGn, 'Participants': SG.n, 'MAE, calibrated, y (95% CI)': [ci(a, b, c, 2) for a, b, c in zip(SG.mae, SG.mae_lo, SG.mae_hi)],
                  'MAE, network output, y': SG.raw_mae.round(2), 'r': SG.r.round(3), 'Mean error, y': SG.bias.round(2)}), '')
B = g('ablation_summary.csv')
SUP['S4'] = ('Model variants: accuracy and DCM contrast',
    pd.DataFrame({'Model': B.label, 'Internal test MAE, y (95% CI)': [ci(a, b, c, 2) for a, b, c in zip(B.test_mae, B.test_mae_lo, B.test_mae_hi)], 'Internal test r': B.test_r.round(3),
                  'External Center 1 MAE, y': B.ext_c1_mae.round(2), 'DCM vs local controls, y (95% CI)': [ci(a, b, c) for a, b, c in zip(B['DCM beta'], B['DCM lo'], B['DCM hi'])],
                  'P': B['DCM p'].map(P_)}),
    'The primary model was selected by a rule fixed before training (internal-test MAE no more than 0.25 y worse and lower scanner/site variance than the base model). Ablations were trained on the base model.')
L = g('loso.csv')
SUP['S5'] = ('Leave-one-source-out validation',
    pd.DataFrame({'Held-out source': L.held_out.map(DS), 'Participants': L.n, 'MAE, y': L.mae.round(2), 'r': L.r.round(3), 'Mean error, y': L.bias.round(2),
                  'Slope of error on age': L.slope.round(3), 'Cross-fitted MAE, same participants, y': L.cv_mae_same_participants.round(2)}),
    'Network output; the held-out source contributed no training data.')
SS = g('contrast_sensitivity.csv'); SS = SS[~SS.analysis.isin(['within Center1', 'within Center2', 'group x center interaction'])]  # center-specific DCM estimates not reported
SUP['S6'] = ('Sensitivity analyses of the pre-specified contrasts',
    pd.DataFrame({'Contrast': SS.contrast.map(CN), 'Analysis': SS.analysis.str.replace('Center1', 'Center 1').str.replace('Center2', 'Center 2').str.replace('C(source)', 'source', regex=False),
                  'Participants': SS.n.astype('Int64'), 'Estimate, y (95% CI)': est(SS), 'P': SS.p.map(P_), 'Note': SS.note.fillna('')}), 'Estimates are differences in calibrated age gap unless stated (interaction terms in years per year or years).')
RR = g('reader_reference.csv'); RR = RR[RR['sample'] == 'controls + abnormal']
PRD = {'n_spinalstenosis': 'Levels with canal stenosis', 'n_foraminalstenosis': 'Levels with foraminal stenosis', 'n_hyperintensity': 'Levels with cord T2 hyperintensity',
       'any cord T2 hyperintensity': 'Any cord T2 hyperintensity', 'exam_umnsigns': 'Upper motor neuron signs', 'exam_motordeficit': 'Motor deficit on examination',
       'exam_gaitabnormal': 'Gait abnormality on examination', 'sx_motor': 'Motor symptoms', 'sx_gaitabnormal': 'Gait symptoms', 'sx_neckpain': 'Neck pain'}
RA = g('reader_agreement.csv')
SUP['S7'] = ('Radiologist-reported and clinical findings, Center 1 (controls and abnormal MRI)',
    pd.DataFrame({'Finding': RR.predictor.map(PRD), 'Unit': RR.unit, 'Participants': RR.n, 'Estimate, y (95% CI)': est(RR), 'P': RR.p.map(P_), 'Partial ρ': RR.partial_rho.round(3)}),
    'Age-adjusted association with calibrated age gap. Agreement between independent codings (κ): ' + '; '.join(f"{r.finding.replace('spinalstenosis', 'canal stenosis').replace('foraminalstenosis', 'foraminal stenosis')} {r.kappa:.2f}" for r in RA.itertuples()) + '.')
FO = g('focal_levels.csv')
SUP['S8'] = ('Within-person focal excess at affected levels',
    pd.DataFrame({'Analysis': FO.analysis, 'Participants': FO.participants, 'Affected vs unaffected levels, y (95% CI)': [ci(a, b, c, 2) for a, b, c in zip(FO.beta, FO.lo, FO.hi)],
                  'P (mixed model)': FO.p.map(P_), 'P (Wilcoxon)': FO.wilcoxon_p.map(P_)}), 'Linear mixed model with a random intercept per participant.')
PR = g('proteomics.csv')
SUP['S9'] = ('Plasma proteins and whole-spine age gap in DCM',
    pd.DataFrame({'Protein': PR.protein, 'Participants': PR.n, 'Partial Spearman ρ': PR.rho.round(3), 'P': PR.p.map(P_), 'FDR q': PR.q.round(3)}),
    'Partial rank correlation adjusted for age, sex, assay batch and center; Benjamini–Hochberg FDR across all proteins.')
AC = g('acquisition.csv'); ALB = ['OpenNeuro vs CSpineSeg', 'Spine Generic vs CSpineSeg', '3D vs 2D acquisition', 'Philips vs GE (Spine Generic)', 'Siemens vs GE (Spine Generic)',
                                  'Slice thickness, per mm (CSpineSeg)', '3 T vs 1.5 T (Center 1 DCM)']
SUP['S10'] = ('Acquisition and data-source associations with age gap', pd.DataFrame({'Comparison': ALB, 'Participants': AC.n, 'Difference, y (95% CI)': est(AC), 'P': AC.p.map(P_)}), 'Healthy participants unless stated.')
RL = g('reliability.csv'); PO = g('posture.csv'); jit = PO[PO.analysis.str.contains('jitter')].iloc[0]
SUP['S11'] = ('Repeat-scan reliability and robustness to landmark perturbation',
    pd.concat([pd.DataFrame({'Analysis': ['Repeat scans, all series', 'Repeat scans, same contrast'], 'Pairs / participants': RL.n, 'ICC(2,1)': RL.icc_2_1.round(3),
                             'Mean difference, y': RL.bias.round(2), '95% limits of agreement, y': [f'{a:.1f} to {b:.1f}' for a, b in zip(RL.loa_lo, RL.loa_hi)]}),
               pd.DataFrame({'Analysis': ['Vertebral anchor jitter (SD 1.5 mm)'], 'Pairs / participants': [int(jit.n)], 'ICC(2,1)': [round(jit.beta, 3)],
                             'Mean difference, y': [round(jit.lo, 2)], '95% limits of agreement, y': ['']})]), '')
MR = g('marrow.csv')
SUP['S12'] = ('Marrow signal and body mass index', pd.DataFrame({'Analysis': MR.analysis.str.replace('g_spine', 'Whole-spine gap').str.replace('g_bone', 'Bone gap').str.replace('g_cord', 'Cord gap')
                              .str.replace('g_csf', 'CSF gap').str.replace('g_disc', 'Disc gap'), 'Participants': MR.n, 'Estimate (95% CI)': est(MR, 2), 'P': MR.p.map(P_)}), '')
HA = g('harmonization_all.csv'); HA = HA[~HA.contrast.isin(['DCM vs controls, Center 1', 'DCM vs controls, Center 2'])]
SUP['S13'] = ('Harmonization and preprocessing sensitivity analyses',
    pd.DataFrame({'Contrast': HA.contrast.map(lambda x: CN.get(x, x)), 'Analysis': HA.analysis.str.replace('vendor_field', 'vendor × field').str.replace('= model', '= scanner model'),
                  'Compartment': HA.marker.fillna(''), 'Controls / cases': [f'{int(a)} / {int(b)}' if pd.notna(a) else '' for a, b in zip(HA.n_ctrl, HA.n_case)],
                  'Estimate (95% CI)': [ci(a, b, c) if pd.notna(b) else (f'{a:.3f}' if a < 1 else f'{a:.2f}') for a, b, c in zip(HA.beta, HA.lo, HA.hi)], 'P': HA.p.map(P_)}),
    'ComBat (empirical Bayes) fitted on healthy participants only and applied to all participants; batch = vendor × field strength or scanner model. Batch ICC from a random-intercept model in healthy participants.')

# conventional MRI measures versus age gap (DCM)
IC = g('imaging_vs_gap_clinical.csv'); IC = IC[IC['sample'] == 'available']; JT = g('imaging_vs_gap_joint.csv')
ICt = pd.DataFrame({'Outcome': IC.outcome, 'MRI measure': IC.predictor, 'Participants': IC.n, 'Scale': IC.scale,
                    'Estimate (95% CI)': [ci(a, b, c, 2) for a, b, c in zip(IC.est, IC.lo, IC.hi)], 'P': IC.p.map(P_)})
ICt = pd.concat([ICt, pd.DataFrame({'Outcome': JT.outcome, 'MRI measure': 'Age gap, adjusted for MSCC, T2 hyperintensity and cord area',
                                    'Participants': JT.n, 'Scale': 'points per SD',
                                    'Estimate (95% CI)': [ci(a, b, c, 2) for a, b, c in zip(JT.gap_est, JT.lo, JT.hi)], 'P': JT.p.map(P_)})], ignore_index=True)
OV = json.load(open(TAB/'imaging_vs_gap_overlap.json'))

# final supplementary order follows first citation in the text
OLD = SUP
SUP = {'S1': OLD['S1'],
       'S2': ('Accuracy of the whole-spine age estimate in healthy participants', T2, 'Network output is the uncalibrated estimate; calibration was fitted on out-of-fold development predictions only. Local offset subtracts the leave-one-out mean error of controls from the same center. 95% CIs from 4,000 participant bootstrap resamples. MAE, mean absolute error.'),
       'S3': OLD['S2'], 'S4': OLD['S3'],
       'S5': ('Pre-specified contrasts in age gap between patients and controls from the same center or dataset', T3, 'Differences in calibrated age gap (years) from ordinary least squares with HC3 standard errors. Holm correction across the four whole-spine contrasts; compartment P values (brackets) are Holm-adjusted within each contrast. The calibration bootstrap refits the calibration in 1,000 resamples of the development set.'),
       'S6': OLD['S6'], 'S7': OLD['S7'], 'S8': OLD['S8'],
       'S9': ('Whole-spine age gap and baseline clinical measures in degenerative cervical myelopathy', T4, 'Estimates per SD of age gap (7.4 years) from models adjusted for age, sex and center: linear models with HC3 standard errors (mJOA, mFI-5, compression, duration), logistic regression (T2 hyperintensity) and proportional-odds regression (severity). Lower mJOA indicates worse function.'),
       'S10': ('Age gap and conventional MRI measures as correlates of neurological function, frailty and severity in DCM', ICt,
               'Models adjusted for age, sex and center. Continuous measures are per SD across patients with DCM, signed so that higher values indicate more compression or a smaller, flatter cord; T2 hyperintensity is present versus absent. Severity estimates are odds ratios for a more severe mJOA category. '
               f"Spearman correlations of the age gap with MSCC, cord area and T2 hyperintensity were {OV['spearman']['MSCC_diameter_AP']:.2f}, {OV['spearman']['area']:.2f} and {OV['spearman']['t2hi']:.2f}."),
       'S11': OLD['S9'], 'S12': OLD['S13'], 'S13': OLD['S4'], 'S14': OLD['S10'], 'S15': OLD['S12'], 'S16': OLD['S11'], 'S17': OLD['S5']}


def rl(df):
    df = df.copy(); df.columns = [relabel(c) for c in df.columns]
    if df.index.dtype == object: nm = df.index.name; df.index = [relabel(i) for i in df.index]; df.index.name = nm
    return df.apply(lambda c: c.map(relabel) if c.dtype == object else c)


LAB = {  # plain-language labels for analysis codes
    'age missing': 'Age missing', 'image label Unknown': 'No dataset label', 'image label not labelled': 'No dataset label', 'included': 'Included',
    'other pathology (Intramedullary tumor)': 'Other pathology (intramedullary tumor)', 'straightening failed: error RuntimeError': 'Straightening failed (processing error)',
    'straightening failed: no_centerline': 'Straightening failed (no cord centerline)', 'age >= 90': 'Age ≥ 90 years',
    'incomplete C2-T1 coverage after straightening': 'Incomplete C2–T1 coverage', 'reader findings incomplete': 'Level findings incomplete',
    'segmentation missing': 'Segmentation missing', 'no clinical record': 'No clinical record',
    'primary': 'Primary analysis', 'Primary': 'Primary analysis', '+ sex': 'Adjusted for sex', 'quadratic age': 'Quadratic age adjustment',
    'unweighted pooling': 'Unweighted averaging of voxel estimates', '+ BMI (where available)': 'Adjusted for body-mass index (where available)',
    '+ native curvature': 'Adjusted for cervical curvature', 'common age support': 'Patients within the control age range',
    'overlapping age range': 'Patients and controls within the overlapping age range', 'group x age interaction (per year)': 'Group × age interaction (per year)',
    'age-matched pairs (within center, no replacement, caliper 3 y)': 'Age-matched pairs (within center, ages within 3 years)',
    'excluding reader age disagreement > 1 y': 'Excluding records whose age differed by >1 year between reviewers',
    'prior spine surgery vs controls (descriptive)': 'Previous cervical surgery vs controls (descriptive)',
    'cases including prior spine surgery': 'Patients including those with previous cervical surgery', 'within Center1': 'JHU only', 'within Center2': 'MCW only',
    'within JHU': 'JHU only', 'within MCW': 'MCW only', 'group x center interaction': 'Group × center interaction',
    'secondary: all healthy controls, + source': 'All healthy controls, adjusted for data source', 'secondary: all healthy controls, + protocol': 'All healthy controls, adjusted for 2D/3D protocol',
    'unharmonized, same participants (vendor × field)': 'Not harmonized (participants in vendor × field-strength batches)', 'ComBat, batch = vendor × field': 'ComBat, batch = vendor × field strength',
    'ComBat compartments': 'ComBat, compartment gaps', 'unharmonized, same participants (model)': 'Not harmonized (participants in scanner-model batches)',
    'same scanner models only (AERA, SKYRA, VERIO), + model': 'Same scanner models only (Siemens Aera, Skyra, Verio), adjusted for model',
    '+ field strength, in-plane resolution, slice thickness': 'Adjusted for field strength, in-plane resolution and slice thickness',
    '+ scanner model, field strength, resolution': 'Adjusted for scanner model, field strength and resolution', '3D converted to 3.3-mm 2D-like slices': '3D scans converted to simulated 3.3-mm slices',
    'Whole-spine gap ~ T1 marrow/disc ratio (per SD) + age': 'Whole-spine gap per SD of T1 marrow-to-disc signal ratio (age-adjusted)',
    'Bone gap ~ T1 marrow/disc ratio (per SD) + age': 'Bone gap per SD of T1 marrow-to-disc signal ratio (age-adjusted)',
    'Cord gap ~ T1 marrow/disc ratio (per SD) + age': 'Cord gap per SD of T1 marrow-to-disc signal ratio (age-adjusted)',
    'CSF gap ~ T1 marrow/disc ratio (per SD) + age': 'CSF gap per SD of T1 marrow-to-disc signal ratio (age-adjusted)',
    'Disc gap ~ T1 marrow/disc ratio (per SD) + age': 'Disc gap per SD of T1 marrow-to-disc signal ratio (age-adjusted)',
    'abnormal vs controls, age (participants with T1)': 'Abnormal MRI vs controls, age-adjusted (participants with T1 images)',
    'abnormal vs controls, age + T1 marrow ratio (participants with T1)': 'Abnormal MRI vs controls, adjusted for age and T1 marrow-to-disc ratio',
    'T1 marrow ratio: abnormal vs controls (SD units)': 'T1 marrow-to-disc ratio, abnormal MRI vs controls (SD units)',
    'healthy: gap ~ BMI (per SD) + age + dataset': 'Healthy participants: gap per SD of body-mass index (adjusted for age and data source)'}


def lab(v):
    if not isinstance(v, str): return v
    v = LAB.get(v, v)
    v = re.sub(r'^cases within control age range (\d+)-(\d+): (\d+)/(\d+)$', r'Patients within control age range (\1–\2 y): \3 of \4', v)
    v = re.sub(r'^overlap (\d+)-(\d+); controls (\d+), cases (\d+)$', r'Overlap \1–\2 y; \3 controls, \4 patients', v)
    v = re.sub(r'^controls (\d+), cases (\d+)$', r'\1 controls, \2 patients', v)
    v = v.replace(': reader spinal stenosis levels', ': levels with canal stenosis').replace(': reader foraminal stenosis levels', ': levels with foraminal stenosis')
    v = v.replace('Spine Generic compression: reported levels', 'Spine Generic compression: compressed levels').replace('Cross-validated development', 'Development (cross-fitted)')
    v = v.replace('Without N4 / acquisition augmentation', 'Base model (no N4 or acquisition augmentation)').replace('Vertebral anchor jitter', 'Vertebral landmark perturbation')
    return v.replace('Duke abnormal', 'CSpineSeg abnormal').replace('Duke controls', 'CSpineSeg controls').replace('Duke', 'CSpineSeg')


SUP['S7'] = ('Level findings and clinical examination, JHU (controls and abnormal MRI)', SUP['S7'][1].fillna(''), SUP['S7'][2])
SUP = {k: (t, v.apply(lambda c: c.map(lab) if c.dtype == object else c), n) for k, (t, v, n) in SUP.items()}
T1 = rl(T1); SUP = {k: (relabel(t), rl(v), relabel(n)) for k, (t, v, n) in SUP.items()}
with pd.ExcelWriter(OUT/'Tables_main.xlsx') as w: T1.to_excel(w, sheet_name='Table 1')
with pd.ExcelWriter(OUT/'Tables_supplementary.xlsx') as w:
    for k, (t, v, n) in SUP.items(): v.to_excel(w, sheet_name=f'Table {k}', index=False)


def docx_tables(path, items, landscape=True):
    doc = Document()
    sec = doc.sections[0]
    if landscape:
        sec.orientation = WD_ORIENT.LANDSCAPE; sec.page_width, sec.page_height = sec.page_height, sec.page_width
    for m in ('left_margin', 'right_margin', 'top_margin', 'bottom_margin'): setattr(sec, m, Mm(15))
    st = doc.styles['Normal']; st.font.name = 'Arial'; st.font.size = Pt(8)
    for title, df, note in items:
        p = doc.add_paragraph(); r = p.add_run(title); r.bold = True; r.font.size = Pt(9)
        df = df.reset_index() if df.index.name else df
        t = doc.add_table(rows=1, cols=len(df.columns)); t.style = 'Table Grid'
        for j, c in enumerate(df.columns):
            t.rows[0].cells[j].text = str(c); t.rows[0].cells[j].paragraphs[0].runs[0].bold = True
        for _, row in df.iterrows():
            cells = t.add_row().cells
            for j, v in enumerate(row): cells[j].text = '' if pd.isna(v) else str(v)
        if note: doc.add_paragraph(note).runs[0].italic = True
        doc.add_page_break()
    doc.save(path)


docx_tables(OUT/'Tables_main.docx', [('Table 1. Participant characteristics', T1, '')])
docx_tables(OUT/'Tables_supplementary.docx', [(f'Supplementary Table {k[1:]}. {t}', v, n) for k, (t, v, n) in SUP.items()])
print(T1.to_string()); print(OTHER); print({k: len(v[1]) for k, v in SUP.items()})
